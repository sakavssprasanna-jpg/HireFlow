import os
import io
import uuid
import hashlib
from typing import List, Optional
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from docx import Document as DocxDocument

from ..domain.schemas import DocumentContent, DocumentSection

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB ceiling
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

class FileValidationError(Exception):
    """Raised when an uploaded file violates safety or format rules."""
    pass

class DocumentParsingError(Exception):
    """Raised when an allowed document cannot be parsed into text."""
    pass

class DocumentParser:
    """Production-grade document extraction service for PDF, DOCX, and TXT files."""

    @staticmethod
    def validate_file(filename: str, content: bytes) -> str:
        """
        Validate file extension, size, and magic byte signatures.
        Returns normalized lowercase extension.
        """
        if not content or len(content.strip()) == 0:
            raise FileValidationError("Uploaded file is empty.")

        if len(content) > MAX_FILE_SIZE_BYTES:
            raise FileValidationError(f"File size ({len(content)} bytes) exceeds 10 MB limit.")

        # Sanitize filename: reject path traversal in original filename parameter
        if not filename or ".." in filename or "/" in filename or "\\" in filename:
            raise FileValidationError("Malformed or path-traversing filename detected.")

        basename = os.path.basename(filename).strip()
        if not basename:
            raise FileValidationError("Filename cannot be empty.")

        ext = os.path.splitext(basename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise FileValidationError(f"Unsupported file extension '{ext}'. Allowed: {', '.join(ALLOWED_EXTENSIONS)}")

        # Magic byte signature sniffing
        if ext == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise FileValidationError("File has .pdf extension but lacks valid PDF header signature.")
        elif ext == ".docx":
            if not content.startswith(b"PK\x03\x04"):
                raise FileValidationError("File has .docx extension but lacks valid ZIP/DOCX container signature.")
        elif ext == ".txt":
            # Reject Windows PE executables or ELF binaries renamed to .txt
            if content.startswith(b"MZ") or content.startswith(b"\x7fELF"):
                raise FileValidationError("Binary executable disguised as text file detected and rejected.")

        return ext

    @classmethod
    def parse_document(cls, filename: str, content: bytes) -> DocumentContent:
        """
        Parse raw document bytes into normalized DocumentContent with section/page provenance.
        """
        ext = cls.validate_file(filename, content)
        doc_id = str(uuid.uuid4())
        file_hash = hashlib.sha256(content).hexdigest()
        warnings: List[str] = []

        if ext == ".pdf":
            return cls._parse_pdf(doc_id, filename, content, file_hash, warnings)
        elif ext == ".docx":
            return cls._parse_docx(doc_id, filename, content, file_hash, warnings)
        elif ext == ".txt":
            return cls._parse_txt(doc_id, filename, content, file_hash, warnings)
        else:
            raise FileValidationError(f"Unsupported extension: {ext}")

    @classmethod
    def _parse_pdf(
        cls, doc_id: str, filename: str, content: bytes, file_hash: str, warnings: List[str]
    ) -> DocumentContent:
        try:
            reader = PdfReader(io.BytesIO(content))
        except PdfReadError as e:
            raise DocumentParsingError(f"Malformed PDF file: {str(e)}")
        except Exception as e:
            raise DocumentParsingError(f"Failed to read PDF container: {str(e)}")

        if reader.is_encrypted:
            try:
                # Attempt empty password decrypt
                decrypted = reader.decrypt("")
                if decrypted == 0:
                    raise DocumentParsingError("PDF is password-protected/encrypted and cannot be processed.")
            except Exception:
                raise DocumentParsingError("PDF is password-protected/encrypted and cannot be processed.")

        sections: List[DocumentSection] = []
        full_text_parts: List[str] = []
        current_offset = 0
        total_pages = len(reader.pages)

        for page_idx, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as e:
                warnings.append(f"Warning on page {page_idx}: {str(e)}")
                page_text = ""

            page_clean = page_text.strip()
            if page_clean:
                start_offset = current_offset
                end_offset = start_offset + len(page_clean)
                sections.append(DocumentSection(
                    section_name=f"Page {page_idx}",
                    page_number=page_idx,
                    text=page_clean,
                    start_offset=start_offset,
                    end_offset=end_offset
                ))
                full_text_parts.append(page_clean)
                current_offset = end_offset + 2  # account for paragraph gap

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentParsingError("PDF document contains no extractable text. (Scanned image PDFs are not supported without OCR).")

        return DocumentContent(
            document_id=doc_id,
            filename=filename,
            file_type="PDF",
            file_hash=file_hash,
            sections=sections,
            full_text=full_text,
            warnings=warnings,
            total_pages=total_pages
        )

    @classmethod
    def _parse_docx(
        cls, doc_id: str, filename: str, content: bytes, file_hash: str, warnings: List[str]
    ) -> DocumentContent:
        try:
            doc = DocxDocument(io.BytesIO(content))
        except Exception as e:
            raise DocumentParsingError(f"Malformed or corrupted DOCX file: {str(e)}")

        sections: List[DocumentSection] = []
        full_text_parts: List[str] = []
        current_offset = 0
        current_heading = "Main Content"

        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            # Detect headings
            if p.style and p.style.name and p.style.name.lower().startswith("heading"):
                current_heading = text

            start_offset = current_offset
            end_offset = start_offset + len(text)
            sections.append(DocumentSection(
                section_name=current_heading,
                page_number=1,
                text=text,
                start_offset=start_offset,
                end_offset=end_offset
            ))
            full_text_parts.append(text)
            current_offset = end_offset + 2

        # Extract tables
        for t_idx, table in enumerate(doc.tables, start=1):
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    table_line = " | ".join(row_texts)
                    start_offset = current_offset
                    end_offset = start_offset + len(table_line)
                    sections.append(DocumentSection(
                        section_name=f"Table {t_idx}",
                        page_number=1,
                        text=table_line,
                        start_offset=start_offset,
                        end_offset=end_offset
                    ))
                    full_text_parts.append(table_line)
                    current_offset = end_offset + 2

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentParsingError("DOCX document contains no extractable text.")

        return DocumentContent(
            document_id=doc_id,
            filename=filename,
            file_type="DOCX",
            file_hash=file_hash,
            sections=sections,
            full_text=full_text,
            warnings=warnings,
            total_pages=1
        )

    @classmethod
    def _parse_txt(
        cls, doc_id: str, filename: str, content: bytes, file_hash: str, warnings: List[str]
    ) -> DocumentContent:
        # Decode handling UTF-8, UTF-8 with BOM, and fallback to latin-1
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
                warnings.append("Decoded with latin-1 fallback due to non-UTF8 encoding.")
            except Exception as e:
                raise DocumentParsingError(f"Failed to decode text file: {str(e)}")

        # Normalize line breaks CRLF -> LF
        normalized_text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not normalized_text:
            raise DocumentParsingError("Text file is empty or contains only whitespace.")

        sections = [DocumentSection(
            section_name="Full Text",
            page_number=1,
            text=normalized_text,
            start_offset=0,
            end_offset=len(normalized_text)
        )]

        return DocumentContent(
            document_id=doc_id,
            filename=filename,
            file_type="TXT",
            file_hash=file_hash,
            sections=sections,
            full_text=normalized_text,
            warnings=warnings,
            total_pages=1
        )
